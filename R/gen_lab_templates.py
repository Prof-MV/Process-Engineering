# -*- coding: utf-8 -*-
"""Generate the FOL document templates the Labs chapter (30-Labs.qmd) refers to.

Run:  python R/gen_lab_templates.py   (from the repo root, or anywhere)

Writes into  labs/templates/  (gitignored, like lab_data/ and quizzes/ — these
are staged on FOL by hand, not committed; this script is the source of record).

    Bill_of_Materials_Template_for_Excel.xlsx   -- Design Documentation
    Gantt_chart_template.xlsx                   -- Design Documentation (see note below)
    Risk_Assessment_Template.docx               -- Design Documentation
    week2_lab.ipynb                             -- ML Week 2
    predictive_maintenance.ipynb                -- ML Week 3 (continues in Weeks 4-5)
    streamlit_app.py + requirements.txt         -- ML Week 5
    Gauge_RR_Bar_Stock_Data_Sheet.xlsx          -- Gauge R&R (via gen_grr_sheet.py; also writes the instructor
                                                   answer-key copy into LabInstructions/ if that folder exists)

Requires: openpyxl, python-docx, nbformat  (pip install openpyxl python-docx nbformat)

NOT generated here (deliberately -- see the Labs chapter Teacher-prep notes):
  - the Facility Design & Simulation lab's AnyLogic demo model (needs AnyLogic; a placeholder .alp would be
    useless -- build/fork the real model once and stage it on FOL).
  - the Ergonomics lab's task video/photo sets (real footage of a workplace task).
  - ML Week 3's train_FD001.txt (the real NASA C-MAPSS dataset -- host the actual
    file, do not fabricate turbofan sensor data).
  - ML Week 5's sample_engine.csv (students build it themselves from the real
    C-MAPSS test set as their Task 1 TODO).
  - the Robotic Work-Cell Design lab's RoboDK reference station file (needs RoboDK).

A NOTE ON THE GANTT TEMPLATE: there is no reliable way to hand-write a real
binary .mpp file without Microsoft Project. Instead this script produces an
.xlsx that Microsoft Project opens natively (File > Open > this file), running
its built-in Import Wizard to map the columns to Task Name / Duration /
Predecessors -- a long-standing, documented Project feature. Open it, confirm
the mapping, then File > Save As > .mpp. Update the Design Documentation materials list /
FOL upload if you'd rather keep the historical "Gantt_chart.mpp" filename --
the workflow is the same either way.
"""

import os
import nbformat as nbf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "labs", "templates")
os.makedirs(OUT, exist_ok=True)

HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
EXAMPLE_FILL = PatternFill("solid", fgColor="FDF3EA")
NOTE_FONT = Font(italic=True, color="7D7D7D")


def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ----------------------------------------------------------------------------
# Design Documentation -- Bill of Materials template
# ----------------------------------------------------------------------------
def make_bom():
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM - Shaft Kit"

    ws["A1"] = "Bill of Materials -- Precitech shaft-kit assembly cell (one kit)"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:H1")
    ws["A2"] = ("Part numbers, descriptions and quantities are pre-filled from "
               "assembly drawing 410-2365 Rev B (and the shaft drawing "
               "205-5478 Rev B). Fill in Unit, Source and Supplier for every "
               "row, including the tray and carton (research a real "
               "packaging supplier for those two -- they are not on the "
               "assembly drawing). Source = In-house (machined at Precitech) "
               "or Purchased.")
    ws["A2"].font = NOTE_FONT
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 30
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")

    headers = ["Item #", "Description", "Part Number / Spec Ref",
              "Qty per Kit", "Unit", "Source\n(In-house / Purchased)",
              "Supplier\n(if purchased)", "Notes"]
    hdr_row = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=hdr_row, column=i, value=h)
    style_header_row(ws, hdr_row, len(headers))
    autosize(ws, [8, 30, 22, 12, 10, 20, 20, 30])

    # Scaffold rows: item 1-6 part numbers/descriptions/qty come straight off
    # assembly drawing 410-2365 Rev B; items 7-8 (packaging) are the two part
    # numbers assigned outside the drawing. Source/Supplier are pre-filled
    # for the six drawing items (per the drawing + "purchased components are
    # from McMaster-Carr") and left blank for the two packaging items, which
    # is genuine student research.
    kit_items = [
        # part_number, description, qty, unit, source, supplier
        ("205-5478",   "Shaft (ANSI 4140)",                               1, "ea", "In-house",  ""),
        ("1199N14",    "Spring-Loaded Rotary Shaft Seal with Wiper Lip",  1, "ea", "Purchased", "McMaster-Carr"),
        ("6677K88",    "Tapered-Roller Bearing with Steel Ring",         2, "ea", "Purchased", "McMaster-Carr"),
        ("91595A179",  "Dowel Pin",                                       1, "ea", "Purchased", "McMaster-Carr"),
        ("98541A440",  "External Retaining Ring",                        1, "ea", "Purchased", "McMaster-Carr"),
        ("1523T76",    "Desiccant Sachet",                                1, "ea", "Purchased", "McMaster-Carr"),
        ("910-1473",   "Moulded Tray",                                    1, "ea", "",          ""),
        ("911-4358",   "Labelled Carton",                                 1, "ea", "",          ""),
    ]
    r = hdr_row + 1
    for i, (pn, desc, qty, unit, source, supplier) in enumerate(kit_items, start=1):
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=desc)
        ws.cell(row=r, column=3, value=pn)
        ws.cell(row=r, column=4, value=qty)
        ws.cell(row=r, column=5, value=unit)
        ws.cell(row=r, column=6, value=source)
        ws.cell(row=r, column=7, value=supplier)
        if not source:
            ws.cell(row=r, column=8, value="Research a real packaging supplier")
        r += 1

    dv = DataValidation(type="list", formula1='"In-house,Purchased"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("F%d:F%d" % (hdr_row + 1, r - 1))

    wb.save(os.path.join(OUT, "Bill_of_Materials_Template_for_Excel.xlsx"))


# ----------------------------------------------------------------------------
# Design Documentation -- Gantt chart template (Excel, importable by MS Project)
# ----------------------------------------------------------------------------
def make_gantt():
    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt Template"

    ws["A1"] = "Project schedule -- standing up the Precitech kitting cell"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:F1")

    headers = ["Task Name", "Duration (days)", "Start", "Finish",
              "% Complete", "Predecessors"]
    hdr_row = 3
    for i, h in enumerate(headers, start=1):
        ws.cell(row=hdr_row, column=i, value=h)
    style_header_row(ws, hdr_row, len(headers))
    autosize(ws, [32, 16, 14, 14, 12, 16])

    # The six standup tasks named in the lab, as a starter chain the student
    # edits. Task numbers below refer to the row's position in this list
    # (1-6), which is what Project's Import Wizard expects in Predecessors.
    tasks = [
        ("Procure racking and fixtures", 5, ""),
        ("Install bearing press", 3, "1"),
        ("Set up inspection station", 2, "1"),
        ("Write the SOP", 2, "2,3"),
        ("Train operators", 2, "4"),
        ("Run a pilot batch", 1, "5"),
    ]
    r = hdr_row + 1
    for i, (name, dur, pred) in enumerate(tasks, start=1):
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=dur)
        ws.cell(row=r, column=5, value=0)
        ws.cell(row=r, column=6, value=pred)
        r += 1

    ws2 = wb.create_sheet("Read me first")
    ws2["A1"] = "How to turn this into a Microsoft Project Gantt chart"
    ws2["A1"].font = Font(bold=True, size=12)
    steps = [
        "1. Open Microsoft Project.",
        "2. File > Open > browse to this .xlsx file.",
        "3. Project's Import Wizard opens. Choose 'New map' (or 'Task list' "
        "template if offered).",
        "4. Map the columns: Task Name -> Name, Duration (days) -> Duration, "
        "Predecessors -> Predecessors, % Complete -> % Complete.",
        "5. Finish the wizard. Set the project Start date (File > Project "
        "Information) so Start/Finish dates and the Gantt bars fill in.",
        "6. Edit durations, add resources, and adjust the six starter tasks "
        "as needed for your own schedule.",
        "7. File > Save As > choose Project (.mpp) to save it as a native "
        "Project file from here on.",
    ]
    for i, s in enumerate(steps, start=3):
        ws2.cell(row=i, column=1, value=s)
    ws2.column_dimensions["A"].width = 100
    for i in range(3, 3 + len(steps)):
        ws2.cell(row=i, column=1).alignment = Alignment(wrap_text=True)

    wb.save(os.path.join(OUT, "Gantt_chart_template.xlsx"))


# ----------------------------------------------------------------------------
# Design Documentation -- HSE risk assessment template (Word)
# ----------------------------------------------------------------------------
def make_hse():
    doc = Document()

    title = doc.add_heading("HSE Risk Assessment", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Precitech Components -- shaft-kit assembly & packing cell")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True

    p = doc.add_paragraph()
    p.add_run(
        "Identify at least five health-and-safety risks in the kitting cell "
        "(for example: manual handling of the blank trolley, pinch points at "
        "the bearing press, sharp edges on retaining rings, repetitive "
        "assembly motion, the carton knife). Complete one row per risk. "
        "Delete the example row before you submit."
    )

    table = doc.add_table(rows=1, cols=7)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    headers = ["Hazard", "Who is harmed & how", "Current controls",
              "Further action required", "Owner", "Deadline", "Done"]
    for cell, text in zip(hdr, headers):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True

    example = table.add_row().cells
    example_vals = [
        "EXAMPLE -- pinch point at the bearing press",
        "Operator's fingers; crush injury while positioning the shaft under "
        "the press",
        "Two-hand control on the press; guard interlock",
        "Add a light curtain; review after 3 months",
        "Cell supervisor",
        "2026-01-15",
        "No",
    ]
    for cell, text in zip(example, example_vals):
        cell.text = text

    for _ in range(6):
        table.add_row()

    for col in table.columns:
        for cell in col.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph()
    doc.add_paragraph(
        "Hierarchy of controls, for the 'Further action required' column "
        "(most to least effective): elimination, substitution, engineering "
        "controls, administrative controls, PPE."
    ).runs[0].italic = True

    doc.save(os.path.join(OUT, "Risk_Assessment_Template.docx"))


# ----------------------------------------------------------------------------
# ML Week 2 -- week2_lab.ipynb
# ----------------------------------------------------------------------------
def make_week2_notebook():
    nb = nbf.v4.new_notebook()
    cells = []

    cells.append(nbf.v4.new_markdown_cell(
        "# Machine Learning Week 2: Manufacturing Data and SPC\n\n"
        "Work through the tasks below in order. Every code cell has a `TODO` "
        "where you (or the LLM you direct) fill in the work. Answer every "
        "markdown prompt in place -- this notebook is what you submit."
    ))

    cells.append(nbf.v4.new_markdown_cell(
        "## Task 1 -- Run your Week 1 code on real data\n\n"
        "Paste your `spc_starter.py` from ML Week 1 into the cell below and run "
        "it against `spc_data.csv`. Expect a spike so large it compresses "
        "the chart, and your 3-sigma detector flagging the wrong rows."
    ))
    cells.append(nbf.v4.new_code_cell(
        "# TODO: paste your ML Week 1 spc_starter.py code here, unchanged, and run it.\n"
        "import pandas as pd\n"
        "import matplotlib.pyplot as plt\n"
        "import numpy as np\n"
    ))
    cells.append(nbf.v4.new_markdown_cell(
        "**Task 1 findings (edit this cell):** describe the spike (approximate "
        "value), whether outlier detection worked, and any error you hit.\n\n"
        "*Your answer here.*"
    ))

    cells.append(nbf.v4.new_markdown_cell(
        "## Task 2 -- Find and fix the data-quality issues\n\n"
        "Every change must be recorded in `quality_log` -- **never silently "
        "delete data**. Problems to find: two missing days (timestamp gap -- "
        "document as a network outage, do **not** interpolate); one "
        "impossible diameter (> 30 mm -- decimal correction); a block of rows "
        "with roughness in inches, not micrometres (convert: x 25400); one "
        "negative cycle time (flag and remove); and a real special-cause "
        "cluster around days 43-47 (**keep it** -- document that it is a "
        "genuine event)."
    ))
    cells.append(nbf.v4.new_code_cell(
        "# Load the raw data\n"
        "df = pd.read_csv(\"spc_data.csv\", parse_dates=[\"timestamp\"])\n\n"
        "# Pattern for quality logging -- add one row per correction\n"
        "quality_log = []\n"
        "quality_log.append({\n"
        "    \"row_index\": None,       # TODO\n"
        "    \"column\":    None,       # TODO\n"
        "    \"original\":  None,       # TODO\n"
        "    \"corrected\": None,       # TODO\n"
        "    \"reason\":    \"\",         # TODO\n"
        "    \"action\":    \"\",         # TODO\n"
        "})\n\n"
        "# TODO 1: find the timestamp gap (missing days). Do NOT interpolate --\n"
        "#         just document it in quality_log.\n\n"
        "# TODO 2: find and fix the impossible diameter (> 30 mm).\n\n"
        "# TODO 3: find and fix the roughness-in-inches block (x 25400).\n\n"
        "# TODO 4: find and remove the negative cycle time.\n\n"
        "# TODO 5: confirm the days 43-47 special-cause cluster is still present\n"
        "#         (it is real -- do not remove it).\n\n"
        "df_clean = df.copy()  # apply your fixes to df_clean\n"
    ))
    cells.append(nbf.v4.new_code_cell(
        "# Post-clean checks -- all three should hold before you continue\n"
        "assert df_clean[\"diameter_mm\"].max() < 30, \"an impossible diameter remains\"\n"
        "assert df_clean[\"cycle_time_sec\"].min() >= 0, \"a negative cycle time remains\"\n"
        "assert len(quality_log) >= 4, \"quality_log should have an entry per fix\"\n"
        "print(\"Post-clean checks passed.\")\n"
    ))

    cells.append(nbf.v4.new_markdown_cell(
        "## Task 3 -- Build an X-bar control chart\n\n"
        "Subgroup = 5 parts/day. Control limits use `sigma / sqrt(n)` -- "
        "**not** the standard deviation of the daily means. Spec limits "
        "USL = 25.050, LSL = 24.950."
    ))
    cells.append(nbf.v4.new_code_cell(
        "USL, LSL = 25.050, 24.950\n"
        "n = 5  # parts per subgroup (per day)\n\n"
        "# TODO: build daily subgroups of diameter_mm from df_clean, compute the\n"
        "#       subgroup means (X-bar) and the individuals' sigma, then the\n"
        "#       control limits: UCL/LCL = X-double-bar +/- 3 * sigma / sqrt(n).\n\n"
        "# TODO: plot the X-bar chart with UCL/LCL and identify out-of-control days.\n\n"
        "# TODO: cross-reference the out-of-control days with the operator and\n"
        "#       tool_age_cycles columns.\n\n"
        "# TODO: compute Cpk = min((USL - mu) / (3 * sigma), (mu - LSL) / (3 * sigma))\n"
        "#       using the INDIVIDUALS' sigma, not the subgroup-mean sigma.\n"
    ))

    cells.append(nbf.v4.new_markdown_cell(
        "## Task 4 -- Reflection\n\n"
        "Answer briefly: (1) how many parts had already been made during the "
        "days 43-47 event before the first out-of-control point appeared? "
        "(2) would a model predicting tomorrow's mean diameter from today's "
        "tool age, operator and the last 5 days of data have given earlier "
        "warning than SPC? (3) what data, not in this CSV, would you want to "
        "build a *predictive* rather than *reactive* model?\n\n"
        "*Your answers here.*"
    ))

    nb["cells"] = cells
    with open(os.path.join(OUT, "week2_lab.ipynb"), "w", encoding="utf-8") as f:
        nbf.write(nb, f)


# ----------------------------------------------------------------------------
# ML Weeks 3-5 -- predictive_maintenance.ipynb (capstone notebook)
# ----------------------------------------------------------------------------
def make_capstone_notebook():
    nb = nbf.v4.new_notebook()
    cells = []

    cells.append(nbf.v4.new_markdown_cell(
        "# Predictive Maintenance Capstone -- ML Weeks 3-5\n\n"
        "NASA C-MAPSS FD001 turbofan run-to-failure data. You build this "
        "notebook across three lab sessions: Week 3 (load & explore), Week 4 "
        "(features, train, evaluate, save the model), Week 5 (deploy -- see "
        "`streamlit_app.py`). Direct an LLM to write the Python; you must be "
        "able to explain every line."
    ))

    cells.append(nbf.v4.new_markdown_cell(
        "## Week 3, Task 1 -- Load and label the data\n\n"
        "26 space-separated columns, no header: unit number, time in cycles, "
        "3 operating settings, 21 sensors, and 2 trailing empty columns to "
        "drop. Expect 20,631 rows, 100 engines."
    ))
    cells.append(nbf.v4.new_code_cell(
        "import pandas as pd\n"
        "import numpy as np\n"
        "import matplotlib.pyplot as plt\n"
        "from sklearn.ensemble import RandomForestRegressor\n"
        "from sklearn.metrics import mean_squared_error, mean_absolute_error\n"
        "import joblib\n\n"
        "index_names = ['unit_number', 'time_in_cycles']\n"
        "setting_names = ['op_setting_1', 'op_setting_2', 'op_setting_3']\n"
        "sensor_names = ['sensor_{}'.format(i) for i in range(1, 22)]\n"
        "col_names = index_names + setting_names + sensor_names + ['drop1', 'drop2']\n\n"
        "train = pd.read_csv(\"train_FD001.txt\", sep=' ', header=None, names=col_names)\n"
        "train = train.drop(columns=['drop1', 'drop2'])\n"
        "print(train.shape)  # expect (20631, 26)\n"
        "train.head()\n"
    ))
    cells.append(nbf.v4.new_code_cell(
        "# Compute the RUL label -- there is no RUL column, you compute it:\n"
        "# For each engine, the max cycle is the failure cycle\n"
        "max_cycles = train.groupby(\"unit_number\")[\"time_in_cycles\"].max()\n"
        "# RUL at any row = that engine's failure cycle minus the current cycle\n"
        "train[\"RUL\"] = train[\"unit_number\"].map(max_cycles) - train[\"time_in_cycles\"]\n\n"
        "assert (train.groupby(\"unit_number\").tail(1)[\"RUL\"] == 0).all()\n"
        "print(\"Engines:\", train['unit_number'].nunique())\n"
    ))

    cells.append(nbf.v4.new_markdown_cell(
        "## Week 3, Task 2 -- Visualise one engine's life\n\n"
        "Plot all 21 sensors for engine #1 over its whole life (a 7x3 grid). "
        "Classify sensors as flat / noise / trend."
    ))
    cells.append(nbf.v4.new_code_cell(
        "engine1 = train[train[\"unit_number\"] == 1]\n"
        "fig, axes = plt.subplots(7, 3, figsize=(14, 18))\n"
        "for ax, s in zip(axes.flat, sensor_names):\n"
        "    ax.plot(engine1[\"time_in_cycles\"], engine1[s])\n"
        "    ax.set_title(s, fontsize=8)\n"
        "plt.tight_layout()\n"
        "plt.show()\n\n"
        "# TODO: classify each sensor as flat / noise / trend based on the plots above.\n"
        "flat_sensors = []    # TODO\n"
        "trend_sensors = []   # TODO\n"
    ))
    cells.append(nbf.v4.new_markdown_cell(
        "*Your flat / noise / trend lists here.*"
    ))

    cells.append(nbf.v4.new_markdown_cell(
        "## Week 3, Task 3 -- Identify constant sensors\n\n"
        "std < 0.001 => effectively constant, carries no information. "
        "Expect about 6 in FD001."
    ))
    cells.append(nbf.v4.new_code_cell(
        "sensor_std = train[sensor_names].std().sort_values()\n"
        "constant_sensors = sensor_std[sensor_std < 0.001].index.tolist()\n"
        "print(constant_sensors)\n"
    ))

    cells.append(nbf.v4.new_markdown_cell(
        "## Week 3, Task 4 -- Correlation with RUL\n\n"
        "For the non-constant sensors, |r| > 0.5 marks the most predictive "
        "raw features."
    ))
    cells.append(nbf.v4.new_code_cell(
        "useful_sensors = [s for s in sensor_names if s not in constant_sensors]\n"
        "corr = train[useful_sensors + [\"RUL\"]].corr()[\"RUL\"].drop(\"RUL\")\n"
        "corr.sort_values(key=abs, ascending=False)\n"
    ))

    cells.append(nbf.v4.new_markdown_cell(
        "## Week 3, Task 5 -- Reflection\n\n"
        "(1) plot RUL for all rows in order -- what shape appears, and what "
        "does that say about random row-wise train/validation splits? (2) "
        "which 3 sensors would you most want in a model? (3) what would you "
        "do before dropping the constant sensors in a real plant?\n\n"
        "*Your answers here.*"
    ))

    # --- Week 4 ---------------------------------------------------
    cells.append(nbf.v4.new_markdown_cell(
        "---\n## Week 4 -- Building the model\n\n"
        "Continue in this same notebook."
    ))
    cells.append(nbf.v4.new_code_cell(
        "# TODO Task 1: rolling mean + std (5-cycle window) per useful sensor,\n"
        "#      GROUPED BY ENGINE so the window resets at each engine boundary.\n"
        "#      Verify: engine #2's first-cycle rolling mean == its raw value.\n"
        "train_fe = train.copy()\n"
    ))
    cells.append(nbf.v4.new_code_cell(
        "# TODO Task 2: clip RUL at 125 and plot the before/after histogram.\n"
        "train_fe[\"RUL_clipped\"] = train_fe[\"RUL\"].clip(upper=125)\n"
    ))
    cells.append(nbf.v4.new_code_cell(
        "# TODO Task 3: engine-wise train/validation split -- hold out 20 engines.\n"
        "#      Do NOT split rows randomly. This assert must pass before you train:\n"
        "# overlap = set(train_data[\"unit_number\"]) & set(val_data[\"unit_number\"])\n"
        "# assert len(overlap) == 0, f\"DATA LEAK: engines in both splits: {overlap}\"\n\n"
        "feature_cols = []  # TODO: exclude unit_number, time_in_cycles, RUL,\n"
        "                    #       RUL_clipped, constant sensors, op_setting_*\n"
    ))
    cells.append(nbf.v4.new_code_cell(
        "# TODO Task 4: train the Random Forest.\n"
        "# model = RandomForestRegressor(n_estimators=100, max_depth=10,\n"
        "#                                n_jobs=-1, random_state=42)\n"
        "# model.fit(X_train, y_train)\n"
    ))
    cells.append(nbf.v4.new_code_cell(
        "# TODO Task 5: evaluate against a naive baseline (y_train.mean() for every row).\n"
        "#      Target: RMSE ~ 20-25 cycles, >= 2x better than naive.\n"
        "#      RMSE < 10 almost always means a data leak -- check Task 3.\n"
    ))
    cells.append(nbf.v4.new_code_cell(
        "# TODO Task 6: plot the top 15 features by mean-decrease-in-impurity and\n"
        "#      compare with the Week 3 correlation analysis.\n"
    ))
    cells.append(nbf.v4.new_code_cell(
        "# Task 7: save the model -- the app needs both files.\n"
        "# joblib.dump(model, \"rul_model.pkl\")\n"
        "# joblib.dump(feature_cols, \"feature_cols.pkl\")\n"
    ))

    # --- Week 5 ----------------------------------------------------
    cells.append(nbf.v4.new_markdown_cell(
        "---\n## Week 5 -- Deploying the app\n\n"
        "The rest of this lab happens outside the notebook: complete the "
        "`TODO`s in `streamlit_app.py`, push `rul_model.pkl` and "
        "`feature_cols.pkl` (saved above) plus the app to a public GitHub "
        "repo, and deploy on Streamlit Community Cloud. See the ML Week 5 FOL "
        "quiz for the full steps."
    ))

    nb["cells"] = cells
    with open(os.path.join(OUT, "predictive_maintenance.ipynb"), "w", encoding="utf-8") as f:
        nbf.write(nb, f)


# ----------------------------------------------------------------------------
# ML Week 5 -- streamlit_app.py + requirements.txt
# ----------------------------------------------------------------------------
STREAMLIT_APP = '''\
"""Precitech predictive-maintenance demo -- ML Week 5 starter template.

Complete every TODO, then:  streamlit run streamlit_app.py
"""

import joblib
import numpy as np
import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt

st.set_page_config(page_title="Precitech RUL Predictor", page_icon=":gear:",
                   layout="centered")

ROLLING_WINDOW = 5  # cycles -- must match what you trained with in Week 4


@st.cache_resource
def load_model():
    model = joblib.load("rul_model.pkl")
    feature_cols = joblib.load("feature_cols.pkl")
    return model, feature_cols


def add_rolling_features(df, sensor_cols, window=ROLLING_WINDOW):
    """Rolling mean/std per sensor for a SINGLE engine's cycle history.

    Mirrors the Week 4 feature engineering: no grouping needed here because
    the uploaded CSV is already one engine's data, in cycle order.
    """
    out = df.copy()
    for s in sensor_cols:
        out[f"{s}_rmean"] = out[s].rolling(window, min_periods=1).mean()
        out[f"{s}_rstd"] = out[s].rolling(window, min_periods=1).std().fillna(0)
    return out


def rul_colour(rul):
    if rul < 30:
        return "red", "Urgent -- schedule maintenance soon"
    if rul < 60:
        return "orange", "Monitor closely"
    return "green", "Healthy"


def main():
    model, feature_cols = load_model()

    st.title("Precitech CNC Cell -- Remaining Useful Life Predictor")
    # TODO: add your name and course to this caption.
    st.caption("TODO: Your Name -- ENGR-3027 -- ML Week 5")

    st.write(
        "Upload the last cycles of one engine's sensor readings (C-MAPSS "
        "FD001 format) to estimate its Remaining Useful Life (RUL), in "
        "cycles."
    )

    col1, col2 = st.columns(2)
    uploaded = col1.file_uploader("Upload engine CSV", type="csv")
    # TODO: wire this button to a bundled sample_engine.csv (last 30 cycles
    # of engine #1 from the C-MAPSS test set, which you build in Task 1).
    use_sample = col2.button("Use sample data")

    df = None
    if use_sample:
        try:
            df = pd.read_csv("sample_engine.csv")
        except FileNotFoundError:
            st.error("sample_engine.csv not found -- finish the Task 1 TODO.")
    elif uploaded is not None:
        try:
            df = pd.read_csv(uploaded)
        except Exception as e:
            st.error(f"Could not read that file as a CSV: {e}")

    if df is None:
        st.info("Upload a CSV or click 'Use sample data' to see a prediction.")
        return

    sensor_cols = [c for c in df.columns if c.startswith("sensor_")]
    missing = [c for c in feature_cols
              if c.replace("_rmean", "").replace("_rstd", "") not in df.columns
              and c not in df.columns]
    if not sensor_cols:
        st.error("This CSV has no sensor_* columns -- check you uploaded the "
                "right file.")
        return

    df = df.dropna(subset=sensor_cols)
    fe = add_rolling_features(df, sensor_cols)

    try:
        X = fe[feature_cols].tail(1)
    except KeyError as e:
        st.error(f"This CSV is missing a column the model expects: {e}")
        return

    rul = float(model.predict(X)[0])
    colour, note = rul_colour(rul)
    st.markdown(
        f"### Predicted RUL: :{colour}[{rul:.0f} cycles]  \\n{note}"
    )

    st.subheader("Sensor trend")
    # TODO: let the user pick a sensor; for now, plot the first useful one.
    s = sensor_cols[0]
    fig, ax = plt.subplots()
    ax.plot(df.index, df[s])
    ax.set_xlabel("cycle")
    ax.set_ylabel(s)
    st.pyplot(fig)

    with st.expander("About this model"):
        st.write(
            "**Training data:** NASA C-MAPSS FD001 simulation "
            "(one fault mode, one operating condition).\\n\\n"
            "**Model:** Random Forest, 100 trees, max depth 10.\\n\\n"
            # TODO: put your ACTUAL Week 4 validation RMSE here.
            "**Validation RMSE:** TODO cycles.\\n\\n"
            # TODO: at least one Precitech-specific "not suitable for" statement.
            "**Not suitable for:** TODO -- e.g. this model has never seen "
            "real Precitech CNC Cell 3 data, only a NASA turbofan "
            "simulation; do not use it to schedule real maintenance."
        )


if __name__ == "__main__":
    main()
'''

REQUIREMENTS_TXT = """\
streamlit==1.39.0
scikit-learn==1.5.2
pandas==2.2.3
numpy==1.26.4
matplotlib==3.9.2
joblib==1.4.2
"""


def make_streamlit_starter():
    with open(os.path.join(OUT, "streamlit_app.py"), "w", encoding="utf-8", newline="\n") as f:
        f.write(STREAMLIT_APP)
    with open(os.path.join(OUT, "requirements.txt"), "w", encoding="utf-8", newline="\n") as f:
        f.write(REQUIREMENTS_TXT)


if __name__ == "__main__":
    make_bom()
    make_gantt()
    make_hse()
    make_week2_notebook()
    make_capstone_notebook()
    make_streamlit_starter()
    import gen_grr_sheet   # sits next to this script
    gen_grr_sheet.main()
    print("Wrote FOL templates to", OUT)
    for fn in sorted(os.listdir(OUT)):
        print(" -", fn)
