# -*- coding: utf-8 -*-
"""Generate the Gauge R&R data-collection workbook (bar stock, calipers, gauge block, caliper swap).

Run:  python R/gen_grr_sheet.py   (from the repo root, or anywhere)

Writes
    labs/templates/Gauge_RR_Bar_Stock_Data_Sheet.xlsx
        -- the blank student workbook (stage on FOL; labs/templates/ is gitignored)
    LabInstructions/Quality_Labs/fixed/Gauge RR Bar Stock - EXAMPLE (instructor).xlsx
        -- the same workbook filled with SIMULATED data (answer key; instructors only;
           written only if that folder exists -- LabInstructions/ is gitignored)

Also called from gen_lab_templates.py.

Study design (individual lab with a swap partner)
-------------------------------------------------
Each student is one operator and has their own serial-numbered caliper. Two
students are paired and SWAP calipers, so the study is a fully crossed
2 operators x 2 calipers x 10 bars x 2 trials design, for two characteristics
(diameter, length):

                         my caliper            partner's caliper
    me                   (Me, Mine)            (Me, Theirs)
    partner              (Partner, Mine)       (Partner, Theirs)

A student types their own four reading columns and copies the partner's four
columns (the partner's sheet columns C:F) into columns G:J of their own sheet.
Crossing operators with calipers is what lets the study tell an operator effect
from a caliper effect (with each operator on their own caliper only, the two are
confounded).

Average-and-range method, 5.15-sigma study variation:
    EV  = R-bar * 5.15/d2                   d2  = 1.128 (2 trials); R-bar = mean of the
                                            four condition average ranges
    AV  = sqrt(max(0, (Xdiff_op  * 5.15/d2*)^2 - EV^2/(n r c)))   operator, d2* = 1.41 (2 levels)
    CV  = sqrt(max(0, (Xdiff_cal * 5.15/d2*)^2 - EV^2/(n r c)))   caliper,  d2* = 1.41 (2 levels)
          n r c = 10 bars x 2 trials x 2 (each operator / caliper mean is over 40 readings)
    PV  = Rp * 5.15/d2*                     d2* = 3.18 (10 parts)
    GRR = sqrt(EV^2 + AV^2 + CV^2);  TV = sqrt(GRR^2 + PV^2);  %GRR = 100 GRR/TV
    %tol = 100 GRR/(USL-LSL);  ndc = INT(1.41 PV/GRR)
    Interaction contrast = |(Me,Mine) - (Me,Theirs) - (Partner,Mine) + (Partner,Theirs)| / 2, compared
    with about 2 x (EV/5.15)/sqrt(n r): if it is large the operator effect depends on the caliper.
The operator and caliper effects rest on only two levels each, so they are crude
estimates -- the sheet and the teacher notes say so.
All formulas are live Excel formulas -- nothing is pasted in as a value.

Requires: openpyxl
"""

import os
import random

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, Series
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_PATH = os.path.join(ROOT, "labs", "templates", "Gauge_RR_Bar_Stock_Data_Sheet.xlsx")
EXAMPLE_DIR = os.path.join(ROOT, "LabInstructions", "Quality_Labs", "fixed")
EXAMPLE_PATH = os.path.join(EXAMPLE_DIR, "Gauge RR Bar Stock - EXAMPLE (instructor).xlsx")

N_PARTS, N_TRIALS = 10, 2
N_STUDENTS = 30            # run-order table covers student numbers 1..30 (the number wraps beyond that)
N_SEQ = 4                  # phase 1 trial 1, phase 1 trial 2, phase 2 trial 1, phase 2 trial 2

# The four measurement conditions, in the order of the reading columns.
#   (label, who measured, which caliper)
CONDS = [
    ("Me, my caliper", "Me", "Mine"),
    ("Me, partner's caliper", "Me", "Theirs"),
    ("Partner, partner's caliper", "Partner", "Theirs"),
    ("Partner, my caliper", "Partner", "Mine"),
]
# sheet layout
FIRST, LAST, AVG = 7, 16, 17          # part rows and the average row on the data sheets
READ_COLS = ["C", "D", "E", "F", "G", "H", "I", "J"]     # 4 conditions x 2 trials
RANGE_COLS = ["K", "L", "M", "N"]
AVG_COLS = ["O", "P", "Q", "R"]
PART_AVG = "S"

# ---- styles -----------------------------------------------------------------
HEAD_FILL = PatternFill("solid", fgColor="1F4E5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="D9E5EA")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # type here
PASTE_FILL = PatternFill("solid", fgColor="FCE4D6")   # copy from the partner's sheet
CALC_FILL = PatternFill("solid", fgColor="EFEFEF")    # formula -- do not type
KEY_FILL = PatternFill("solid", fgColor="E2EFDA")
NOTE_FONT = Font(italic=True, color="666666")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def put(ws, ref, value=None, fill=None, font=None, fmt=None, align=None, border=True):
    c = ws[ref]
    if value is not None:
        c.value = value
    if fill is not None:
        c.fill = fill
    if font is not None:
        c.font = font
    if fmt is not None:
        c.number_format = fmt
    if align is not None:
        c.alignment = align
    if border:
        c.border = BOX
    return c


def header(ws, ref, text, span_to=None):
    c = put(ws, ref, text, HEAD_FILL, HEAD_FONT, align=CENTER)
    if span_to:
        ws.merge_cells("%s:%s" % (ref, span_to))
    return c


def title(ws, text, sub=None):
    ws["A1"].value = text
    ws["A1"].font = Font(bold=True, size=14, color="1F4E5F")
    if sub:
        ws["A2"].value = sub
        ws["A2"].font = NOTE_FONT


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ---- Read me ----------------------------------------------------------------
def sheet_readme(wb):
    ws = wb.active
    ws.title = "Read me"
    title(ws, "Gauge R&R with calipers - bar stock (diameter and length), caliper swap (inches)",
          "Yellow cells: you type. Orange cells: you copy from your partner's sheet. Grey cells: formulas - do not type over them.")
    steps = [
        ("1  Setup", "Enter your name and caliper serial number, your partner's, your student number (it picks your "
                     "run order) and the stamped numbers of your 10 bars - in ASCENDING order, so your partner's rows "
                     "line up with yours. Your instructor gives the tolerances; check them."),
        ("2  Caliper check", "With your partner, zero-check BOTH calipers with the jaws closed, then measure the gauge "
                             "block three times with each. The sheet works out each caliper's bias."),
        ("3  Phase 1: own caliper", "Measure all 10 bars with YOUR caliper: trial 1 in the order shown on the Run order "
                                    "sheet, then trial 2 in its (different) order. Measure the diameter AND the length "
                                    "each time you pick a bar up. Type the readings in the Diameter and Length sheets "
                                    "(columns C-D). Do not look at your earlier readings, or at your partner's."),
        ("4  Phase 2: swap", "Swap calipers with your partner and measure all 10 bars again, twice, with THEIR caliper "
                             "(columns E-F). Use the phase 2 orders on the Run order sheet."),
        ("5  Exchange readings", "When both of you have finished, copy your partner's columns C:F into your columns G:J "
                                 "(Diameter and Length). Their C:D are their readings with their own caliper; their E:F "
                                 "are their readings with your caliper."),
        ("6  Read the results", "Everything else calculates: EV, operator effect AV, caliper effect CV, GRR, PV, TV, %GRR, "
                                "%tolerance, ndc and the verdicts. The Summary sheet sets both characteristics side by side."),
    ]
    header(ws, "A4", "Step")
    header(ws, "B4", "What to do")
    r = 5
    for name, txt in steps:
        put(ws, "A%d" % r, name, SUB_FILL, BOLD, align=LEFT)
        put(ws, "B%d" % r, txt, align=LEFT)
        ws.row_dimensions[r].height = 64
        r += 1
    put(ws, "A12", "Method", SUB_FILL, BOLD, align=LEFT)
    put(ws, "B12", "Average-and-range Gauge R&R, 5.15-sigma study variation (99 % of the spread). Two operators (you and your "
                   "partner) each measure with both calipers, so an operator difference and a caliper difference can be told "
                   "apart. Each of those two effects rests on only two levels, so treat the split as an estimate, not a "
                   "verdict. Repeatability EV is the scatter of repeat readings; AV is the operator effect (holding the "
                   "caliper constant); CV is the caliper effect (holding the operator constant).", align=LEFT)
    ws.row_dimensions[12].height = 92
    put(ws, "A13", "Acceptance", SUB_FILL, BOLD, align=LEFT)
    put(ws, "B13", "%GRR of total variation: under 10 % acceptable; 10-30 % marginal; over 30 % unacceptable. "
                   "ndc (number of distinct categories) should be 5 or more.", align=LEFT)
    ws.row_dimensions[13].height = 34
    put(ws, "A15", "Colour key", SUB_FILL, BOLD, align=LEFT)
    put(ws, "B15", "Yellow = type here", INPUT_FILL, align=LEFT)
    put(ws, "B16", "Orange = copy from your partner's sheet", PASTE_FILL, align=LEFT)
    put(ws, "B17", "Grey = calculated by a formula", CALC_FILL, align=LEFT)
    put(ws, "B18", "Green = a Gauge R&R result you report", KEY_FILL, align=LEFT)
    widths(ws, {"A": 24, "B": 112})
    landscape(ws)


# ---- Setup ------------------------------------------------------------------
# Fixed cell addresses other sheets refer to.
ME_NAME, ME_SERIAL = "Setup!$B$7", "Setup!$C$7"
PT_NAME, PT_SERIAL = "Setup!$B$8", "Setup!$C$8"
STUDENT_NO = "Setup!$B$9"
BAR_RANGE = "Setup!$B$12:$B$21"
RES = "Setup!$B$24"
TOL = {"Diameter": ("Setup!$D$26", "Setup!$E$26"), "Length": ("Setup!$D$27", "Setup!$E$27")}
K1, K2, K3 = "Setup!$B$36", "Setup!$B$37", "Setup!$B$38"
NPARTS, NTRIALS, NREAD = "Setup!$B$31", "Setup!$B$30", "Setup!$B$40"
MULT = "Setup!$B$39"


def sheet_setup(wb, ex=None):
    ws = wb.create_sheet("Setup")
    title(ws, "Setup", "Fill in the yellow cells before you start measuring.")
    put(ws, "A3", "Section / lab group", SUB_FILL, BOLD)
    put(ws, "B3", ex["team"] if ex else None, INPUT_FILL)
    ws.merge_cells("B3:C3")
    put(ws, "A4", "Date", SUB_FILL, BOLD)
    put(ws, "B4", ex["date"] if ex else None, INPUT_FILL)
    ws.merge_cells("B4:C4")

    header(ws, "A6", "Person")
    header(ws, "B6", "Name")
    header(ws, "C6", "Caliper serial no.")
    put(ws, "A7", "Me", SUB_FILL, BOLD)
    put(ws, "B7", ex["me"][0] if ex else None, INPUT_FILL)
    put(ws, "C7", ex["me"][1] if ex else None, INPUT_FILL)
    put(ws, "A8", "My swap partner", SUB_FILL, BOLD)
    put(ws, "B8", ex["partner"][0] if ex else None, INPUT_FILL)
    put(ws, "C8", ex["partner"][1] if ex else None, INPUT_FILL)
    put(ws, "A9", "My student number (1-30; picks my run order)", SUB_FILL, BOLD)
    put(ws, "B9", ex["student_no"] if ex else None, INPUT_FILL, align=CENTER)
    put(ws, "C9", "your partner must have a different number", font=NOTE_FONT, border=False)

    header(ws, "A11", "Study part")
    header(ws, "B11", "Bar number (stamped) - ascending order")
    for i in range(N_PARTS):
        r = 12 + i
        put(ws, "A%d" % r, i + 1, SUB_FILL, BOLD, align=CENTER)
        put(ws, "B%d" % r, ex["bars"][i] if ex else None, INPUT_FILL, align=CENTER)

    header(ws, "A23", "Gauge and specification", "E23")
    put(ws, "A24", "Caliper resolution (in)", SUB_FILL, BOLD)
    put(ws, "B24", 0.0005, INPUT_FILL, fmt="0.0000")
    header(ws, "A25", "Characteristic")
    header(ws, "B25", "Nominal (in)")
    header(ws, "C25", "Tolerance +/- (in)")
    header(ws, "D25", "LSL (in)")
    header(ws, "E25", "USL (in)")
    # The nominals/tolerances are PLACEHOLDERS -- the instructor replaces them with the bars' real values.
    for r, name, nom, tol in ((26, "Diameter", 1.0, 0.005), (27, "Length", 2.0, 0.020)):
        put(ws, "A%d" % r, name, SUB_FILL, BOLD)
        put(ws, "B%d" % r, nom, INPUT_FILL, fmt="0.0000")
        put(ws, "C%d" % r, tol, INPUT_FILL, fmt="0.0000")
        put(ws, "D%d" % r, "=B%d-C%d" % (r, r), CALC_FILL, fmt="0.0000")
        put(ws, "E%d" % r, "=B%d+C%d" % (r, r), CALC_FILL, fmt="0.0000")
    put(ws, "A28", "Instructor: the nominal and tolerance values above are placeholders - replace them "
                   "with the specification for your bars.", font=NOTE_FONT, border=False)

    header(ws, "A29", "Constants (change only if the study design changes)", "E29")
    consts = [
        (30, "Trials per condition, r", 2, "0"),
        (31, "Parts, n", 10, "0"),
        (32, "Calipers (and operators), levels of each", 2, "0"),
        (33, "d2 for the average range (r = 2 trials)", 1.128, "0.000"),
        (34, "d2* for the operator and caliper averages (2 levels)", 1.41, "0.00"),
        (35, "d2* for the part averages (n = 10)", 3.18, "0.00"),
    ]
    for r, label, val, fmt in consts:
        put(ws, "A%d" % r, label, SUB_FILL, BOLD)
        put(ws, "B%d" % r, val, INPUT_FILL, fmt=fmt)
    put(ws, "A36", "K1 = 5.15 / d2 (repeatability)", SUB_FILL, BOLD)
    put(ws, "B36", "=$B$39/B33", CALC_FILL, fmt="0.00000")
    put(ws, "A37", "K2 = 5.15 / d2* (operator and caliper effects)", SUB_FILL, BOLD)
    put(ws, "B37", "=$B$39/B34", CALC_FILL, fmt="0.00000")
    put(ws, "A38", "K3 = 5.15 / d2* (parts)", SUB_FILL, BOLD)
    put(ws, "B38", "=$B$39/B35", CALC_FILL, fmt="0.00000")
    put(ws, "A39", "Study-variation multiplier (5.15 sigma = 99 %)", SUB_FILL, BOLD)
    put(ws, "B39", 5.15, INPUT_FILL, fmt="0.00")
    put(ws, "A40", "Readings behind each operator (or caliper) average, n r x levels", SUB_FILL, BOLD)
    put(ws, "B40", "=B31*B30*B32", CALC_FILL, fmt="0")
    widths(ws, {"A": 58, "B": 22, "C": 22, "D": 14, "E": 14})
    landscape(ws)


# ---- Caliper check ----------------------------------------------------------
def sheet_caliper(wb, ex=None):
    ws = wb.create_sheet("Caliper check")
    title(ws, "Caliper check against the gauge block",
          "Do this before the study, with your partner: check both calipers. Clean the jaws and the block; handle the block by its edges only.")
    put(ws, "A3", "Gauge block serial no.", SUB_FILL, BOLD)
    put(ws, "B3", ex["block_serial"] if ex else None, INPUT_FILL)
    put(ws, "A4", "Certified size (in)", SUB_FILL, BOLD)
    put(ws, "B4", ex["block_size"] if ex else None, INPUT_FILL, fmt="0.00000")
    put(ws, "A5", "Acceptable bias, +/- (in)", SUB_FILL, BOLD)
    put(ws, "B5", "=2*%s" % RES, CALC_FILL, fmt="0.0000")
    put(ws, "C5", "two resolution steps; your instructor may change this", font=NOTE_FONT, border=False)

    heads = ["Caliper", "Serial no.", "Zero check, jaws closed (in)", "Block reading 1 (in)",
             "Block reading 2 (in)", "Block reading 3 (in)", "Mean of block readings (in)",
             "Bias = mean - certified (in)", "Range of the 3 readings (in)", "Zero OK?", "Bias within limit?"]
    for j, h in enumerate(heads):
        header(ws, "%s7" % get_column_letter(j + 1), h)
    ws.row_dimensions[7].height = 46
    for i, (label, serial) in enumerate((("My caliper", ME_SERIAL), ("Partner's caliper", PT_SERIAL))):
        r = 8 + i
        put(ws, "A%d" % r, label, SUB_FILL, BOLD)
        put(ws, "B%d" % r, "=%s" % serial, CALC_FILL)
        for j in range(4):
            col = get_column_letter(3 + j)
            v = ex["caliper"][i][j] if ex else None
            put(ws, "%s%d" % (col, r), v, INPUT_FILL, fmt="0.0000")
        put(ws, "G%d" % r, '=IF(COUNT(D%d:F%d)=3,AVERAGE(D%d:F%d),"")' % (r, r, r, r), CALC_FILL, fmt="0.0000")
        put(ws, "H%d" % r, '=IF(AND(ISNUMBER(G%d),ISNUMBER($B$4)),G%d-$B$4,"")' % (r, r), CALC_FILL, fmt="+0.0000;-0.0000;0.0000")
        put(ws, "I%d" % r, '=IF(COUNT(D%d:F%d)=3,MAX(D%d:F%d)-MIN(D%d:F%d),"")' % ((r,) * 6), CALC_FILL, fmt="0.0000")
        put(ws, "J%d" % r, '=IF(ISNUMBER(C%d),IF(ABS(C%d)<=%s,"OK","Re-zero"),"")' % (r, r, RES), CALC_FILL, align=CENTER)
        put(ws, "K%d" % r, '=IF(ISNUMBER(H%d),IF(ROUND(ABS(H%d),8)<=ROUND($B$5,8),"OK","Too large - tell instructor"),"")' % (r, r), CALC_FILL, align=CENTER)
    put(ws, "A10", "Difference in bias, my caliper - partner's caliper (in)", SUB_FILL, BOLD, align=LEFT)
    ws.merge_cells("A10:G10")
    put(ws, "H10", '=IF(AND(ISNUMBER(H8),ISNUMBER(H9)),H8-H9,"")', KEY_FILL, fmt="+0.0000;-0.0000;0.0000")
    put(ws, "A12", "A single block checks a caliper at one size only. It says nothing about the caliper's "
                   "linearity over its range, or about the length of your bars.", font=NOTE_FONT, border=False)
    widths(ws, {"A": 26, "B": 22, "C": 18, "D": 14, "E": 14, "F": 14, "G": 16, "H": 18, "I": 16, "J": 11, "K": 26})
    landscape(ws)


# ---- Run order --------------------------------------------------------------
def run_orders(seed=3027):
    """One random permutation of the 10 parts per student number x sequence (fixed seed: reproducible file).

    Sequences: 1 = phase 1 trial 1, 2 = phase 1 trial 2, 3 = phase 2 trial 1, 4 = phase 2 trial 2.
    """
    rng = random.Random(seed)
    table = []
    for s in range(1, N_STUDENTS + 1):
        for q in range(1, N_SEQ + 1):
            p = list(range(1, N_PARTS + 1))
            rng.shuffle(p)
            table.append((s, q, p))
    return table


ORDER_RANGE = "'Order table'!$C$5:$L$%d" % (4 + N_STUDENTS * N_SEQ)


def sheet_order(wb):
    ws = wb.create_sheet("Run order")
    title(ws, "Measurement order (randomised - it depends on your student number on the Setup sheet)",
          "Entries are your bars' stamped numbers ('Part n' until you have entered them). Measure BOTH the diameter and "
          "the length of each bar in the order shown.")
    header(ws, "A4", "Position")
    labels = ["Phase 1 (my caliper)\nTrial 1", "Phase 1 (my caliper)\nTrial 2",
              "Phase 2 (partner's caliper)\nTrial 1", "Phase 2 (partner's caliper)\nTrial 2"]
    for q in range(N_SEQ):
        col = get_column_letter(2 + q)
        header(ws, "%s4" % col, labels[q])
        for pos in range(N_PARTS):
            k = "INDEX(%s,MOD(%s-1,%d)*%d+%d,%d)" % (ORDER_RANGE, STUDENT_NO, N_STUDENTS, N_SEQ, q + 1, pos + 1)
            put(ws, "%s%d" % (col, 5 + pos),
                '=IF(ISNUMBER(%s),IF(INDEX(%s,%s)="","Part "&%s,INDEX(%s,%s)),"enter student no.")' %
                (STUDENT_NO, BAR_RANGE, k, k, BAR_RANGE, k), CALC_FILL, align=CENTER)
    ws.row_dimensions[4].height = 46
    for pos in range(N_PARTS):
        put(ws, "A%d" % (5 + pos), pos + 1, SUB_FILL, BOLD, align=CENTER)
    put(ws, "A16", "Finish trial 1 for all bars before starting trial 2, and do not look up your earlier readings. "
                   "Do not look at your partner's readings until both of you have finished.",
        font=NOTE_FONT, border=False)
    widths(ws, {"A": 10, "B": 24, "C": 24, "D": 26, "E": 26})
    landscape(ws)


def sheet_order_table(wb):
    ws = wb.create_sheet("Order table")
    title(ws, "Order table (do not edit)",
          "Random orders of the 10 study parts for student numbers 1-%d; the Run order sheet looks up yours." % N_STUDENTS)
    header(ws, "A4", "Student no.")
    header(ws, "B4", "Sequence")
    for j in range(N_PARTS):
        header(ws, "%s4" % get_column_letter(3 + j), "Pos %d" % (j + 1))
    for i, (s, q, p) in enumerate(run_orders()):
        r = 5 + i
        put(ws, "A%d" % r, s, SUB_FILL, align=CENTER)
        put(ws, "B%d" % r, q, SUB_FILL, align=CENTER)
        for j, v in enumerate(p):
            put(ws, "%s%d" % (get_column_letter(3 + j), r), v, align=CENTER)
    widths(ws, {"A": 12, "B": 10})
    ws.sheet_properties.tabColor = "A6A6A6"


# ---- Data + results sheets (Diameter, Length) --------------------------------
RES_ROW = {}   # key -> row on the data sheets (identical on both)

EXPECTED_ROWS = {"Rbar": 21, "EV": 22, "mMe": 23, "mPt": 24, "XdOp": 25, "AV": 26, "mMine": 27, "mTheirs": 28,
                 "XdCal": 29, "CV": 30, "GRR": 31, "Rp": 32, "PV": 33, "TV": 34,
                 "pEV": 36, "pAV": 37, "pCV": 38, "pGRR": 39, "pPV": 40, "pTol": 41, "ndc": 42,
                 "Int": 44, "IntNoise": 45, "IntV": 46,
                 "v1": 48, "v2": 49, "v3": 50, "big": 51}


def sheet_data(wb, name, ex=None):
    ws = wb.create_sheet(name)
    title(ws, "%s - readings and Gauge R&R (inches)" % name,
          "Yellow: your readings. Orange: copy your partner's columns C:F of THEIR %s sheet here (into G:J). Everything else calculates." % name)
    # ---- header rows 4-6
    header(ws, "A4", "Part")
    header(ws, "B4", "Bar no.")
    ws.merge_cells("A4:A6")
    ws.merge_cells("B4:B6")
    header(ws, "C4", '="Readings I took ("&%s&")"' % ME_NAME, "F4")
    header(ws, "G4", '="Readings my partner took ("&%s&")"' % PT_NAME, "J4")
    header(ws, "C5", '="My caliper ("&%s&")"' % ME_SERIAL, "D5")
    header(ws, "E5", '="Partner\'s caliper ("&%s&")"' % PT_SERIAL, "F5")
    header(ws, "G5", '="Partner\'s caliper ("&%s&")"' % PT_SERIAL, "H5")
    header(ws, "I5", '="My caliper ("&%s&")"' % ME_SERIAL, "J5")
    for k, col in enumerate(READ_COLS):
        header(ws, "%s6" % col, "Trial %d" % (1 + k % 2))
    header(ws, "K4", "Range of the 2 trials", "N4")
    header(ws, "O4", "Average of the 2 trials", "R4")
    for k, (label, _o, _c) in enumerate(CONDS):
        for grp in (RANGE_COLS, AVG_COLS):
            header(ws, "%s5" % grp[k], label)
            ws.merge_cells("%s5:%s6" % (grp[k], grp[k]))
    header(ws, "%s4" % PART_AVG, "Part average")
    ws.merge_cells("%s4:%s6" % (PART_AVG, PART_AVG))
    ws.row_dimensions[5].height = 34
    ws.row_dimensions[6].height = 20

    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="100",
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Reading", error="Enter a number in inches, e.g. 1.0005")
    ws.add_data_validation(dv)
    for i in range(N_PARTS):
        r = FIRST + i
        put(ws, "A%d" % r, i + 1, SUB_FILL, BOLD, align=CENTER)
        put(ws, "B%d" % r, '=IF(INDEX(%s,A%d)="","",INDEX(%s,A%d))' % (BAR_RANGE, r, BAR_RANGE, r), CALC_FILL, align=CENTER)
        for k, col in enumerate(READ_COLS):
            v = ex["data"][name][k // 2][k % 2][i] if ex else None
            put(ws, "%s%d" % (col, r), v, INPUT_FILL if k < 4 else PASTE_FILL, fmt="0.0000", align=CENTER)
            dv.add("%s%d" % (col, r))
        for c in range(4):
            a, b = READ_COLS[2 * c], READ_COLS[2 * c + 1]
            put(ws, "%s%d" % (RANGE_COLS[c], r),
                '=IF(COUNT(%s%d:%s%d)=2,MAX(%s%d:%s%d)-MIN(%s%d:%s%d),"")' % ((a, r, b, r) * 3),
                CALC_FILL, fmt="0.0000", align=CENTER)
            put(ws, "%s%d" % (AVG_COLS[c], r),
                '=IF(COUNT(%s%d:%s%d)=2,AVERAGE(%s%d:%s%d),"")' % ((a, r, b, r) * 2),
                CALC_FILL, fmt="0.00000", align=CENTER)
        put(ws, "%s%d" % (PART_AVG, r), '=IF(COUNT(O%d:R%d)=4,AVERAGE(O%d:R%d),"")' % (r, r, r, r), CALC_FILL, fmt="0.00000", align=CENTER)
    put(ws, "A%d" % AVG, "Average", SUB_FILL, BOLD, align=CENTER)
    ws.merge_cells("A%d:J%d" % (AVG, AVG))
    for col in RANGE_COLS + AVG_COLS:
        put(ws, "%s%d" % (col, AVG),
            '=IF(COUNT(%s%d:%s%d)=10,AVERAGE(%s%d:%s%d),"")' % (col, FIRST, col, LAST, col, FIRST, col, LAST),
            KEY_FILL, BOLD, "0.00000", CENTER)
    put(ws, "%s%d" % (PART_AVG, AVG), '=IF(COUNT(S%d:S%d)=10,AVERAGE(S%d:S%d),"")' % (FIRST, LAST, FIRST, LAST),
        KEY_FILL, BOLD, "0.00000", CENTER)

    # ---- results block -------------------------------------------------------------
    header(ws, "A19", "Gauge R&R results - study variation = 5.15 sigma (in)", "S19")
    lo, hi = TOL[name]
    O, P, Q, R_ = ("%s%d" % (c, AVG) for c in AVG_COLS)      # cell means: (Me,Mine) (Me,Theirs) (Partner,Theirs) (Partner,Mine)
    rows = [
        ("Rbar", "Average range, R-bar (mean of the four conditions)", '=IF(COUNT(K17:N17)=4,AVERAGE(K17:N17),"")', "0.00000", "average of the four conditions' average ranges"),
        ("EV", "Repeatability, EV", '=IF(ISNUMBER(F21),%s*F21,"")' % K1, "0.00000", "EV = K1 x R-bar"),
        ("mMe", "Operator: my average (both calipers)", '=IF(AND(ISNUMBER(%s),ISNUMBER(%s)),AVERAGE(%s,%s),"")' % (O, P, O, P), "0.00000", "mean of (me, my caliper) and (me, partner's caliper)"),
        ("mPt", "Operator: partner's average (both calipers)", '=IF(AND(ISNUMBER(%s),ISNUMBER(%s)),AVERAGE(%s,%s),"")' % (Q, R_, Q, R_), "0.00000", "mean of (partner, partner's caliper) and (partner, my caliper)"),
        ("XdOp", "Operator difference, X-diff (operator)", '=IF(AND(ISNUMBER(F23),ISNUMBER(F24)),ABS(F23-F24),"")', "0.00000", "|my average - partner's average|; the caliper is held constant in this comparison"),
        ("AV", "Operator effect, AV", '=IF(AND(ISNUMBER(F22),ISNUMBER(F25)),SQRT(MAX(0,(F25*%s)^2-F22^2/%s)),"")' % (K2, NREAD), "0.00000", "AV = sqrt[ (X-diff x K2)^2 - EV^2 / (n r x 2) ]; 0 if the root is negative"),
        ("mMine", "Caliper: my caliper's average (both operators)", '=IF(AND(ISNUMBER(%s),ISNUMBER(%s)),AVERAGE(%s,%s),"")' % (O, R_, O, R_), "0.00000", "mean of (me, my caliper) and (partner, my caliper)"),
        ("mTheirs", "Caliper: partner's caliper's average (both operators)", '=IF(AND(ISNUMBER(%s),ISNUMBER(%s)),AVERAGE(%s,%s),"")' % (P, Q, P, Q), "0.00000", "mean of (me, partner's caliper) and (partner, partner's caliper)"),
        ("XdCal", "Caliper difference, X-diff (caliper)", '=IF(AND(ISNUMBER(F27),ISNUMBER(F28)),ABS(F27-F28),"")', "0.00000", "|my caliper - partner's caliper|; the operator is held constant in this comparison"),
        ("CV", "Caliper effect, CV", '=IF(AND(ISNUMBER(F22),ISNUMBER(F29)),SQRT(MAX(0,(F29*%s)^2-F22^2/%s)),"")' % (K2, NREAD), "0.00000", "CV = sqrt[ (X-diff x K2)^2 - EV^2 / (n r x 2) ]; 0 if the root is negative"),
        ("GRR", "Gauge R&R, GRR", '=IF(AND(ISNUMBER(F22),ISNUMBER(F26),ISNUMBER(F30)),SQRT(F22^2+F26^2+F30^2),"")', "0.00000", "GRR = sqrt(EV^2 + AV^2 + CV^2)"),
        ("Rp", "Range of part averages, Rp", '=IF(COUNT(S7:S16)=10,MAX(S7:S16)-MIN(S7:S16),"")', "0.00000", "largest minus smallest part average (all four conditions)"),
        ("PV", "Part variation, PV", '=IF(ISNUMBER(F32),%s*F32,"")' % K3, "0.00000", "PV = K3 x Rp"),
        ("TV", "Total variation, TV", '=IF(AND(ISNUMBER(F31),ISNUMBER(F33)),SQRT(F31^2+F33^2),"")', "0.00000", "TV = sqrt(GRR^2 + PV^2)"),
        None,
        ("pEV", "%EV of total variation", '=IF(ISNUMBER(F34),100*F22/F34,"")', "0.0", "100 x EV / TV"),
        ("pAV", "%AV (operator) of total variation", '=IF(ISNUMBER(F34),100*F26/F34,"")', "0.0", "100 x AV / TV"),
        ("pCV", "%CV (caliper) of total variation", '=IF(ISNUMBER(F34),100*F30/F34,"")', "0.0", "100 x CV / TV"),
        ("pGRR", "%GRR of total variation", '=IF(ISNUMBER(F34),100*F31/F34,"")', "0.0", "100 x GRR / TV  (the number the acceptance bands apply to)"),
        ("pPV", "%PV of total variation", '=IF(ISNUMBER(F34),100*F33/F34,"")', "0.0", "100 x PV / TV"),
        ("pTol", "%GRR of tolerance", '=IF(AND(ISNUMBER(F31),ISNUMBER(%s),ISNUMBER(%s)),100*F31/(%s-%s),"")' % (hi, lo, hi, lo), "0.0", "100 x GRR / (USL - LSL)"),
        ("ndc", "Number of distinct categories, ndc", '=IF(AND(ISNUMBER(F33),ISNUMBER(F31)),IF(F31=0,"",INT(1.41*F33/F31)),"")', "0", "ndc = INT( 1.41 x PV / GRR ); 5 or more is adequate"),
        None,
        ("Int", "Operator-by-caliper interaction contrast", '=IF(AND(ISNUMBER(%s),ISNUMBER(%s),ISNUMBER(%s),ISNUMBER(%s)),ABS((%s-%s-%s+%s)/2),"")' % (O, P, Q, R_, O, P, R_, Q), "0.00000", "|(me,mine) - (me,theirs) - (partner,mine) + (partner,theirs)| / 2"),
        ("IntNoise", "Interaction noise level (about 2 sigma)", '=IF(ISNUMBER(F22),2*(F22/%s)/SQRT(%s*%s),"")' % (MULT, NPARTS, NTRIALS), "0.00000", "2 x (EV / 5.15) / sqrt(n r): how big the contrast could be from repeat scatter alone"),
        ("IntV", "Interaction verdict", '=IF(AND(ISNUMBER(F44),ISNUMBER(F45)),IF(F44>F45,"Noticeable: the operator effect depends on the caliper","Within noise: operator and caliper effects add up"),"")', None, ""),
        None,
        ("v1", "Verdict on %GRR of total variation", '=IF(ISNUMBER(F39),IF(F39<10,"Acceptable (under 10 %)",IF(F39<=30,"Marginal (10 to 30 %)","Unacceptable (over 30 %)")),"")', None, ""),
        ("v2", "Verdict on %GRR of tolerance", '=IF(ISNUMBER(F41),IF(F41<10,"Acceptable (under 10 %)",IF(F41<=30,"Marginal (10 to 30 %)","Unacceptable (over 30 %)")),"")', None, ""),
        ("v3", "Verdict on ndc", '=IF(ISNUMBER(F42),IF(F42>=5,"Adequate (5 or more)","Not adequate (under 5)"),"")', None, ""),
        ("big", "Largest component of the gauge variation", '=IF(AND(ISNUMBER(F22),ISNUMBER(F26),ISNUMBER(F30)),IF(AND(F22>=F26,F22>=F30),"Repeatability (EV)",IF(F26>=F30,"Operator (AV)","Caliper (CV)")),"")', None, ""),
    ]
    key_rows = ("EV", "AV", "CV", "GRR", "PV", "TV", "pGRR", "pTol", "ndc")
    r = 21
    for row in rows:
        if row is None:
            r += 1
            continue
        key, label, formula, fmt, expl = row
        put(ws, "A%d" % r, label, SUB_FILL, BOLD, align=LEFT)
        ws.merge_cells("A%d:E%d" % (r, r))
        is_text = key in ("v1", "v2", "v3", "big", "IntV")
        put(ws, "F%d" % r, formula, KEY_FILL if (key in key_rows or is_text) else CALC_FILL,
            BOLD if is_text else None, fmt, LEFT if is_text else CENTER)
        if is_text:
            ws.merge_cells("F%d:K%d" % (r, r))
        else:
            put(ws, "G%d" % r, expl, font=NOTE_FONT, border=False)
        RES_ROW[key] = r
        r += 1
    assert RES_ROW == EXPECTED_ROWS, RES_ROW     # the hard-coded F-references above rely on these rows

    red = PatternFill("solid", bgColor="F8CBAD")
    grn = PatternFill("solid", bgColor="C6E0B4")
    amb = PatternFill("solid", bgColor="FFE699")
    for rr in (48, 49):
        ws.conditional_formatting.add("F%d:K%d" % (rr, rr), FormulaRule(formula=['LEFT($F%d,4)="Unac"' % rr], fill=red))
        ws.conditional_formatting.add("F%d:K%d" % (rr, rr), FormulaRule(formula=['LEFT($F%d,4)="Marg"' % rr], fill=amb))
        ws.conditional_formatting.add("F%d:K%d" % (rr, rr), FormulaRule(formula=['LEFT($F%d,4)="Acce"' % rr], fill=grn))
    ws.conditional_formatting.add("F50:K50", FormulaRule(formula=['LEFT($F50,3)="Not"'], fill=red))
    ws.conditional_formatting.add("F50:K50", FormulaRule(formula=['LEFT($F50,3)="Ade"'], fill=grn))

    # ---- chart: part averages for the four conditions ------------------------------
    ch = LineChart()
    ch.title = "%s: part averages by operator and caliper" % name
    ch.y_axis.title = "Average reading (in)"
    ch.x_axis.title = "Part"
    ch.height, ch.width = 9, 19
    for k, (label, _o, _c) in enumerate(CONDS):
        col = 15 + k                      # O..R
        s = Series(Reference(ws, min_col=col, min_row=FIRST, max_row=LAST), title=label)
        s.marker.symbol = "circle"
        s.smooth = False
        ch.series.append(s)
    ch.set_categories(Reference(ws, min_col=1, min_row=FIRST, max_row=LAST))
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ws.add_chart(ch, "L21")
    widths(ws, {"A": 8, "B": 9, "C": 11, "D": 11, "E": 11, "F": 14, "G": 11, "H": 11, "I": 11, "J": 11,
                "K": 12, "L": 12, "M": 12, "N": 12, "O": 12, "P": 12, "Q": 12, "R": 12, "S": 12})
    ws.freeze_panes = "C7"
    landscape(ws)
    return ws


# ---- Summary ------------------------------------------------------------------
def sheet_summary(wb):
    ws = wb.create_sheet("Summary")
    title(ws, "Summary - diameter and length", "All values come from the Diameter and Length sheets.")
    header(ws, "A4", "Result (5.15-sigma study variation, in)")
    header(ws, "B4", "Diameter")
    header(ws, "C4", "Length")
    items = [("Repeatability, EV", "EV", "0.00000"), ("Operator effect, AV", "AV", "0.00000"),
             ("Caliper effect, CV", "CV", "0.00000"), ("Gauge R&R, GRR", "GRR", "0.00000"),
             ("Part variation, PV", "PV", "0.00000"), ("Total variation, TV", "TV", "0.00000"),
             ("%EV", "pEV", "0.0"), ("%AV (operator)", "pAV", "0.0"), ("%CV (caliper)", "pCV", "0.0"),
             ("%GRR of total variation", "pGRR", "0.0"), ("%PV", "pPV", "0.0"),
             ("%GRR of tolerance", "pTol", "0.0"), ("ndc", "ndc", "0"),
             ("Largest component", "big", None), ("Interaction", "IntV", None),
             ("Verdict (total variation)", "v1", None), ("Verdict (tolerance)", "v2", None), ("Verdict (ndc)", "v3", None)]
    r = 5
    for label, key, fmt in items:
        put(ws, "A%d" % r, label, SUB_FILL, BOLD, align=LEFT)
        for col, sh in (("B", "Diameter"), ("C", "Length")):
            put(ws, "%s%d" % (col, r), "=%s!F%d" % (sh, RES_ROW[key]), KEY_FILL, None, fmt, CENTER)
        ws.row_dimensions[r].height = 30 if key in ("IntV", "big") else 18
        r += 1
    r += 1
    header(ws, "A%d" % r, "Operator or caliper? Signed differences (in)", "D%d" % r)
    r += 1
    for j, h in enumerate(["", "Diameter", "Length", "Gauge block"]):
        header(ws, "%s%d" % (get_column_letter(1 + j), r), h)
    r += 1
    put(ws, "A%d" % r, "Caliper difference: my caliper - partner's caliper (from the bars)", SUB_FILL, BOLD, align=LEFT)
    put(ws, "B%d" % r, '=IF(AND(ISNUMBER(Diameter!F27),ISNUMBER(Diameter!F28)),Diameter!F27-Diameter!F28,"")', CALC_FILL, fmt="+0.0000;-0.0000;0.0000", align=CENTER)
    put(ws, "C%d" % r, '=IF(AND(ISNUMBER(Length!F27),ISNUMBER(Length!F28)),Length!F27-Length!F28,"")', CALC_FILL, fmt="+0.0000;-0.0000;0.0000", align=CENTER)
    put(ws, "D%d" % r, "='Caliper check'!H10", CALC_FILL, fmt="+0.0000;-0.0000;0.0000", align=CENTER)
    ws.row_dimensions[r].height = 32
    r += 1
    put(ws, "A%d" % r, "Operator difference: me - partner (from the bars)", SUB_FILL, BOLD, align=LEFT)
    put(ws, "B%d" % r, '=IF(AND(ISNUMBER(Diameter!F23),ISNUMBER(Diameter!F24)),Diameter!F23-Diameter!F24,"")', CALC_FILL, fmt="+0.0000;-0.0000;0.0000", align=CENTER)
    put(ws, "C%d" % r, '=IF(AND(ISNUMBER(Length!F23),ISNUMBER(Length!F24)),Length!F23-Length!F24,"")', CALC_FILL, fmt="+0.0000;-0.0000;0.0000", align=CENTER)
    put(ws, "D%d" % r, "n/a", CALC_FILL, align=CENTER)
    ws.row_dimensions[r].height = 32
    put(ws, "A%d" % (r + 2), "If the caliper difference measured on the bars matches the difference in gauge-block bias, the block explains the "
                             "caliper effect. The block checks one size only, so a mismatch does not mean the block is wrong. "
                             "The operator difference has no block equivalent: it is how you and your partner handle the "
                             "caliper and the bar (where the jaws sit, how hard they are closed).",
        font=NOTE_FONT, border=False)
    ws.merge_cells("A%d:D%d" % (r + 2, r + 2))
    ws["A%d" % (r + 2)].alignment = LEFT
    ws.row_dimensions[r + 2].height = 62
    widths(ws, {"A": 62, "B": 26, "C": 26, "D": 16})
    landscape(ws)


# ---- example (simulated) data ---------------------------------------------------
def simulate(seed=2027):
    """Simulated readings for the instructor's answer-key copy. NOT measured data."""
    rng = random.Random(seed)
    block = 1.0000
    res = 0.0005

    def rd(x):
        return round(round(x / res) * res, 4)

    cal_bias = [0.0000, -0.0007]            # my caliper, partner's caliper (in)
    op_off = {"Diameter": [0.0000, 0.0002], "Length": [0.0000, 0.0015]}   # me, partner (technique, e.g. where the jaws sit)
    part_sd = {"Diameter": 0.0008, "Length": 0.014}
    nom = {"Diameter": 1.0008, "Length": 2.0000}
    rep_sd = {"Diameter": 0.0005, "Length": 0.0012}
    # conditions in reading-column order: (operator, caliper)
    cond_oc = [(0, 0), (0, 1), (1, 1), (1, 0)]
    data = {}
    for ch in ("Diameter", "Length"):
        true = [rng.gauss(nom[ch], part_sd[ch]) for _ in range(N_PARTS)]
        data[ch] = [[[rd(true[i] + op_off[ch][o] + cal_bias[c] + rng.gauss(0, rep_sd[ch])) for i in range(N_PARTS)]
                     for _t in range(N_TRIALS)] for (o, c) in cond_oc]
    caliper = [[rd(rng.gauss(0, 0.0002))] + [rd(block + cal_bias[c] + rng.gauss(0, 0.0002)) for _ in range(3)]
               for c in range(2)]
    return dict(team="Section 2 (SIMULATED example)", date="example", me=("Student 1", "SN-1041"),
                partner=("Student 2", "SN-1057"), student_no=7, bars=[3, 7, 12, 14, 18, 21, 22, 25, 27, 30],
                block_serial="GB-1000-0417", block_size=block, caliper=caliper, data=data)


# ---- build ------------------------------------------------------------------------
def build(path, example=False):
    ex = simulate() if example else None
    wb = Workbook()
    sheet_readme(wb)
    sheet_setup(wb, ex)
    sheet_caliper(wb, ex)
    sheet_order(wb)
    sheet_data(wb, "Diameter", ex)
    sheet_data(wb, "Length", ex)
    sheet_summary(wb)
    sheet_order_table(wb)
    if example:
        wb["Read me"]["A20"].value = ("INSTRUCTOR COPY - every reading in this workbook is SIMULATED, not measured. "
                                      "Use it as an answer key for the formulas; do not give it to students.")
        wb["Read me"]["A20"].font = Font(bold=True, color="C00000")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path


def main():
    print("wrote", build(TEMPLATE_PATH))
    if os.path.isdir(EXAMPLE_DIR):
        print("wrote", build(EXAMPLE_PATH, example=True))
    else:
        print("skipped the instructor example (no", EXAMPLE_DIR + ")")


if __name__ == "__main__":
    main()
